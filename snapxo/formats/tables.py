# CSV and ODS from the standard library alone, XLSX through openpyxl, which is
# what makes its charts real Excel charts you can still edit.

import csv
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

FORMATS = ("xlsx", "ods", "csv")


def safe_name(text: str) -> str:
    kept = "".join(c if c.isalnum() or c in " -_" else "-" for c in text).strip()
    return "-".join(kept.lower().split()) or "sheet"


def file_name(base: str, extension: str, stamp: str) -> str:
    return f"snapxo-{safe_name(base)}-{stamp}.{extension}"


def write_csv(dataset: dict, target: Path) -> Path:
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(dataset["columns"])
        writer.writerows(dataset["rows"])
    return target


def write_csv_folder(datasets: list[dict], folder: Path, stamp: str) -> list[Path]:
    # One file per dataset, because CSV holds exactly one table.
    folder.mkdir(parents=True, exist_ok=True)
    return [write_csv(dataset, folder / file_name(dataset["key"], "csv", stamp))
            for dataset in datasets]


def write_xlsx(datasets: list[dict], target: Path, with_charts: bool = True) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    book = Workbook()
    book.remove(book.active)

    for dataset in datasets:
        sheet = book.create_sheet(title=_sheet_title(dataset["title"]))
        sheet.append(dataset["columns"])
        for row in dataset["rows"]:
            sheet.append(list(row))

        if not with_charts or not dataset.get("chart") or len(dataset["rows"]) < 2:
            continue

        chart = {"line": LineChart, "bar": BarChart, "donut": PieChart}[dataset["chart"]]()
        chart.title = dataset["title"]
        last_row = len(dataset["rows"]) + 1
        last_column = len(dataset["columns"])
        chart.add_data(Reference(sheet, min_col=2, max_col=last_column, min_row=1, max_row=last_row),
                       titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=last_row))
        chart.height, chart.width = 8, 20
        sheet.add_chart(chart, f"{chr(ord('A') + last_column + 1)}2")

    book.save(target)
    return target


def _sheet_title(title: str) -> str:
    # Excel refuses these characters and anything past 31 letters.
    cleaned = "".join(c for c in title if c not in "[]:*?/\\")
    return cleaned[:31] or "Sheet"


def write_ods(datasets: list[dict], target: Path) -> Path:
    sheets = "".join(_ods_table(dataset) for dataset in datasets)
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'office:version="1.3"><office:body><office:spreadsheet>'
        f"{sheets}</office:spreadsheet></office:body></office:document-content>"
    )
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        'manifest:version="1.3">'
        '<manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )

    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        # The mimetype has to be the first entry and stored, not deflated.
        archive.writestr(zipfile.ZipInfo("mimetype"),
                         "application/vnd.oasis.opendocument.spreadsheet",
                         compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/manifest.xml", manifest)
        archive.writestr("content.xml", content)
    return target


def _ods_table(dataset: dict) -> str:
    rows = [_ods_row(dataset["columns"])]
    rows += [_ods_row(row) for row in dataset["rows"]]
    name = escape(_sheet_title(dataset["title"]), {'"': "&quot;"})
    columns = f'<table:table-column table:number-columns-repeated="{len(dataset["columns"])}"/>'
    return f'<table:table table:name="{name}">{columns}{"".join(rows)}</table:table>'


def _ods_row(values) -> str:
    return f'<table:table-row>{"".join(_ods_cell(value) for value in values)}</table:table-row>'


def _ods_cell(value) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (f'<table:table-cell office:value-type="float" office:value="{value}">'
                f"<text:p>{value}</text:p></table:table-cell>")
    return (f'<table:table-cell office:value-type="string">'
            f"<text:p>{escape(str(value))}</text:p></table:table-cell>")
